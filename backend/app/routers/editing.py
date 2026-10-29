"""Chat-with-the-form edits (as tracked changes), propagation to sibling ICFs, EC response table."""

import io
import json
import logging
import tempfile
from collections.abc import Iterator
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from pathlib import Path
from typing import Literal
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from .. import db, services, storage
from ..auth import User, current_staff
from .projects import _doc_row, _project_row, _version_row

log = logging.getLogger("icf.editing")
router = APIRouter(prefix="/api", dependencies=[Depends(current_staff)], tags=["editing"])

GUARDED_TYPES = ("adult", "parent")  # forms that follow the national template verbatim


# ---------- helpers ----------


@contextmanager
def _docx_path(v) -> Iterator[str]:
    """Local filesystem path of a version's .docx (downloaded from GCS when needed)."""
    if not v["filename"].lower().endswith(".docx"):
        raise HTTPException(422, "Editing is only available for .docx documents")
    try:
        cm = storage.local_path(v["stored_name"])
        path = cm.__enter__()
    except FileNotFoundError:
        raise HTTPException(404, "File missing from storage") from None
    try:
        yield str(path)
    finally:
        cm.__exit__(None, None, None)


def _require_ai() -> None:
    from ..engine import ai

    if not ai.ai_status()["enabled"]:
        raise HTTPException(503, "Vertex AI is not available (check gcloud auth / project)")


def _template_key(language: str | None) -> str | None:
    """Storage key of the newest national template in the library for this language."""
    with db.connect() as conn:
        row = conn.execute(
            "SELECT stored_name FROM documents WHERE kind = 'template' AND language = ?"
            " AND filename LIKE '%.docx' ORDER BY COALESCE(version, '') DESC, id DESC LIMIT 1",
            (language or "en",),
        ).fetchone()
    return row["stored_name"] if row else None


def _annotate(edits: list[dict], paragraphs: list[dict], icf_type, language) -> list[dict]:
    """Add section + mandatory-template match to each proposed edit."""
    from ..engine import edit

    by_idx = {p["idx"]: p for p in paragraphs}
    mandatory: tuple[str, ...] = ()
    if icf_type in GUARDED_TYPES:
        key = _template_key(language)
        if key:
            try:
                with storage.local_path(key) as tpl:
                    mandatory = edit.template_mandatory(str(tpl))
            except Exception:
                log.exception("could not read template %s", key)
    for e in edits:
        e["section"] = by_idx.get(e["idx"], {}).get("section", "")
        e["mandatory"] = (
            edit.mandatory_match(edit.plain(e["old"]), mandatory) if mandatory else None
        )
    return edits


def _latest_version(conn, doc_id: int):
    return conn.execute(
        "SELECT * FROM versions WHERE doc_id = ? ORDER BY number DESC LIMIT 1", (doc_id,)
    ).fetchone()


def _write_version(conn_vid: int, edits: list[dict], note: str, source: str, user: dict):
    """Apply edits to version conn_vid's file -> (new version id, skipped). No AI inside txn."""
    from ..engine import edit

    with db.connect() as conn:
        v = _version_row(conn, conn_vid)
        doc = _doc_row(conn, v["doc_id"])
    with _docx_path(v) as src, tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "edited.docx"
        result = edit.apply_edits(src, str(out), edits, author=user["name"] or edit.AUTHOR)
        data = out.read_bytes()
    if not result["applied"]:
        raise HTTPException(409, "No edit could be applied (the document changed?)")
    items = services.prepare_comments(v["filename"], data)
    with db.connect() as conn:
        from ..engine import trace

        base = trace.version_source(conn, v["id"], "edited from" if source == "edited" else source)
        vid = services.add_version(
            conn, doc["id"], v["filename"], data, note, source, user["name"], items,
            sources=[base] if base else None,
        )  # fmt: skip
        services.carry_decisions(conn, v["id"], vid)
    return vid, result["skipped"]


# ---------- paragraphs + chat ----------


@router.get("/versions/{vid}/paragraphs")
def paragraphs(vid: int) -> list[dict]:
    from ..engine import edit

    with db.connect() as conn:
        v = _version_row(conn, vid)
    with _docx_path(v) as path:
        return edit.list_paragraphs(path)


class ChatMsg(BaseModel):
    role: Literal["user", "assistant"]
    text: str


class ChatIn(BaseModel):
    message: str
    history: list[ChatMsg] = Field(default_factory=list)
    comment_id: int | None = None


@router.post("/versions/{vid}/chat")
def chat(vid: int, body: ChatIn) -> dict:
    from ..engine import assist, edit

    if not body.message.strip():
        raise HTTPException(422, "message is required")
    _require_ai()
    with db.connect() as conn:
        v = _version_row(conn, vid)
        doc = _doc_row(conn, v["doc_id"])
        comment = None
        if body.comment_id:
            row = conn.execute("SELECT * FROM comments WHERE id = ?", (body.comment_id,))
            comment = row.fetchone()
    with _docx_path(v) as path:
        paras = edit.list_paragraphs(path)
    result = assist.propose(
        paras,
        instruction=body.message.strip(),
        history=[m.model_dump() for m in body.history],
        icf_type=doc["icf_type"],
        language=doc["language"],
        comment=dict(comment) if comment else None,
    )
    if result is None:
        raise HTTPException(502, "Vertex AI did not return a usable answer, please retry")
    _annotate(result["edits"], paras, doc["icf_type"], doc["language"])
    return result


class EditIn(BaseModel):
    idx: int
    old: str
    new: str
    reason: str = ""


class ApplyIn(BaseModel):
    edits: list[EditIn]
    instruction: str = ""
    comment_id: int | None = None


@router.post("/versions/{vid}/apply")
def apply(vid: int, body: ApplyIn, user: User) -> dict:
    if not body.edits:
        raise HTTPException(422, "No edits to apply")
    edits = [e.model_dump() for e in body.edits]
    note = f"Edited: {body.instruction.strip()[:200]}" if body.instruction.strip() else "Edited"
    new_vid, skipped = _write_version(vid, edits, note, "edited", user)
    with db.connect() as conn:
        v = _version_row(conn, vid)
        doc = _doc_row(conn, v["doc_id"])
        new_v = _version_row(conn, new_vid)
        if body.comment_id:
            src = conn.execute(
                "SELECT author, text FROM comments WHERE id = ?", (body.comment_id,)
            ).fetchone()
            if src:
                note_txt = f"Addressed in v{new_v['number']}"
                for vv in (vid, new_vid):
                    conn.execute(
                        "UPDATE comments SET status = 'approved', decided_by = ?, decided_at = ?,"
                        " decision_note = ? WHERE version_id = ? AND author = ? AND text = ?",
                        (user["name"], db.now_iso(), note_txt, vv, src["author"], src["text"]),
                    )
        cur = conn.execute(
            "INSERT INTO changes (project_id, doc_id, source_version_id, result_version_id,"
            " instruction, icf_type, language, edits_json, created_by, created_at)"
            " VALUES (?,?,?,?,?,?,?,?,?,?)",
            (doc["project_id"], doc["id"], vid, new_vid, body.instruction.strip(),
             doc["icf_type"], doc["language"], json.dumps(edits, ensure_ascii=False),
             user["name"], db.now_iso()),
        )  # fmt: skip
        return {
            "version": services.version_dict(conn, new_v),
            "change_id": cur.lastrowid,
            "skipped": skipped,
        }


# ---------- propagation ----------


def _change_row(conn, cid: int):
    row = conn.execute("SELECT * FROM changes WHERE id = ?", (cid,)).fetchone()
    if not row:
        raise HTTPException(404, "Change not found")
    return row


@router.get("/changes/{cid}")
def get_change(cid: int) -> dict:
    with db.connect() as conn:
        c = _change_row(conn, cid)
        d = dict(c)
        d["edits"] = json.loads(d.pop("edits_json"))
        return d


def _preview_targets(change: dict, project_id: int, doc_id: int) -> list[dict]:
    from ..engine import assist, edit

    _require_ai()
    with db.connect() as conn:
        rows = conn.execute(
            "SELECT * FROM project_docs WHERE project_id = ? AND role = 'icf' AND id != ?"
            " ORDER BY id",
            (project_id, doc_id),
        ).fetchall()
        targets = []
        for d in rows:
            v = _latest_version(conn, d["id"])
            if v and v["filename"].lower().endswith(".docx"):
                targets.append((services.doc_dict(conn, d, with_versions=False), v))

    def one(target):
        doc, v = target
        base = {"doc": doc, "version_id": v["id"], "version_number": v["number"]}
        try:
            with _docx_path(v) as path:
                paras = edit.list_paragraphs(path)
            res = assist.propagate(change, paras, icf_type=doc["icf_type"],
                                   language=doc["language"])  # fmt: skip
        except Exception as exc:  # noqa: BLE001
            log.exception("propagate preview failed")
            return {**base, "status": "error", "reason": str(exc), "edits": []}
        if res is None:
            return {**base, "status": "error", "reason": "Vertex AI call failed", "edits": []}
        _annotate(res["edits"], paras, doc["icf_type"], doc["language"])
        status = "proposed" if res["applicable"] else "not_applicable"
        return {**base, "status": status, "reason": res["reason"], "edits": res["edits"]}

    with ThreadPoolExecutor(6) as pool:
        return list(pool.map(one, targets))


def _combined_change(conn, cids: list[int]):
    """Merge several changes of ONE doc (oldest first) into one change for the AI."""
    rows = [_change_row(conn, c) for c in dict.fromkeys(cids)]
    if not rows:
        raise HTTPException(422, "No changes selected")
    if len({r["doc_id"] for r in rows}) != 1:
        raise HTTPException(422, "All changes must belong to the same form")
    rows.sort(key=lambda r: r["id"])
    first = rows[0]
    edits = [e for r in rows for e in json.loads(r["edits_json"])]
    instr = "; ".join(r["instruction"] for r in rows if r["instruction"])
    change = {
        "instruction": instr,
        "icf_type": first["icf_type"],
        "language": first["language"],
        "edits": edits,
    }
    return rows, change


@router.post("/changes/{cid}/propagate/preview")
def propagate_preview(cid: int) -> list[dict]:
    with db.connect() as conn:
        c = _change_row(conn, cid)
    change = {
        "instruction": c["instruction"],
        "icf_type": c["icf_type"],
        "language": c["language"],
        "edits": json.loads(c["edits_json"]),
    }
    return _preview_targets(change, c["project_id"], c["doc_id"])


class BatchIn(BaseModel):
    change_ids: list[int]


@router.post("/changes/propagate/preview")
def propagate_preview_batch(body: BatchIn) -> list[dict]:
    """One AI call per target form for all selected changes combined."""
    with db.connect() as conn:
        rows, change = _combined_change(conn, body.change_ids)
    return _preview_targets(change, rows[0]["project_id"], rows[0]["doc_id"])


class TargetIn(BaseModel):
    version_id: int
    edits: list[EditIn]


class PropagateIn(BaseModel):
    targets: list[TargetIn]


class PropagateBatchIn(PropagateIn):
    change_ids: list[int]


def _propagate_write(rows, instruction: str, targets: list[TargetIn], user) -> list[dict]:
    with db.connect() as conn:
        src = _version_row(conn, rows[-1]["result_version_id"])
        src_doc = _doc_row(conn, src["doc_id"])
    label = f"{(src_doc['language'] or '').upper()} {src_doc['icf_type'] or ''} v{src['number']}"
    note = f"Propagated from {label.strip()}: {instruction[:160]}".rstrip(": ")
    out = []
    for t in targets:
        if not t.edits:
            continue
        try:
            vid, skipped = _write_version(
                t.version_id, [e.model_dump() for e in t.edits], note, "propagated", user
            )
        except HTTPException as exc:
            out.append({"version_id": t.version_id, "error": exc.detail})
            continue
        with db.connect() as conn:
            out.append({
                "version_id": t.version_id,
                "version": services.version_dict(conn, _version_row(conn, vid)),
                "skipped": skipped,
            })  # fmt: skip
    if any("version" in o for o in out):
        with db.connect() as conn:
            conn.execute(
                f"UPDATE changes SET propagated_at = ? WHERE id IN ({','.join('?' * len(rows))})",  # noqa: S608
                (db.now_iso(), *[r["id"] for r in rows]),
            )
    return out


@router.post("/changes/{cid}/propagate/apply")
def propagate_apply(cid: int, body: PropagateIn, user: User) -> list[dict]:
    with db.connect() as conn:
        c = _change_row(conn, cid)
    return _propagate_write([c], c["instruction"], body.targets, user)


@router.post("/changes/propagate/apply")
def propagate_apply_batch(body: PropagateBatchIn, user: User) -> list[dict]:
    """Apply the combined edits once per target form (one new version each)."""
    with db.connect() as conn:
        rows, change = _combined_change(conn, body.change_ids)
    return _propagate_write(rows, change["instruction"], body.targets, user)


@router.get("/versions/{vid}/changes")
def pending_changes(vid: int) -> list[dict]:
    """Applied edits of this form (any version) not yet propagated, oldest first."""
    with db.connect() as conn:
        v = _version_row(conn, vid)
        rows = conn.execute(
            "SELECT * FROM changes WHERE doc_id = ? AND propagated_at IS NULL ORDER BY id",
            (v["doc_id"],),
        ).fetchall()
        return [
            {
                "id": r["id"],
                "instruction": r["instruction"],
                "edit_count": len(json.loads(r["edits_json"])),
                "result_version_id": r["result_version_id"],
                "created_by": r["created_by"],
                "created_at": r["created_at"],
            }
            for r in rows
        ]


@router.post("/changes/{cid}/dismiss")
def dismiss_change(cid: int) -> dict:
    with db.connect() as conn:
        _change_row(conn, cid)
        conn.execute("UPDATE changes SET propagated_at = ? WHERE id = ?", (db.now_iso(), cid))
    return {"ok": True}


# ---------- EC response table ----------

STATUS_LABEL = {"open": "Open", "approved": "Accepted", "declined": "Declined (rebuttal)"}


@router.get("/versions/{vid}/ec-response/file")
def ec_response(vid: int, user: User, lang: str | None = None):
    from docx import Document
    from docx.shared import Pt

    lang = services.check_lang(lang) or user["language"]
    with db.connect() as conn:
        v = _version_row(conn, vid)
        doc = _doc_row(conn, v["doc_id"])
        project = _project_row(conn, doc["project_id"])
        items = [c for c in services.comments_for_version(conn, vid, lang) if not c["parent_id"]]
    out = Document()
    out.add_heading("Response to Ethics Committee comments", level=1)
    out.add_paragraph(
        f"{project['name']} – {doc['title']} (v{v['number']}, {v['filename']}). "
        f"Generated {db.now_iso()[:10]} by {user['name']}."
    )
    table = out.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for cell, h in zip(
        table.rows[0].cells,
        ("#", "EC comment", f"Translation ({lang.upper()})", "Decision", "Response / change"),
        strict=True,
    ):
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
    for i, c in enumerate(items, 1):
        row = table.add_row().cells
        row[0].text = str(i)
        row[1].text = f"{c['text']}\n\n[{c['author']}] on: “{c['anchor_text'][:200]}”"
        row[2].text = c["translated_text"] or "—"
        row[3].text = STATUS_LABEL.get(c["status"], c["status"])
        row[4].text = c["decision_note"] or ""
    for r in table.rows:
        for cell in r.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    buf = io.BytesIO()
    out.save(buf)
    buf.seek(0)
    name = f"EC response - {doc['title']} v{v['number']}.docx".replace("/", "-")
    return StreamingResponse(
        buf,
        media_type=services.DOCX_MIME,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"},
    )
