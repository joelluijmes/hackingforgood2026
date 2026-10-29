"""Paragraph-level editing of a docx as Word tracked changes, plus the mandatory-text guard.

Paragraph index = position of the w:p in `body.iter(w:p)` (tables included), so the same
index addresses the same paragraph when listing and when applying edits to that file.
"""

import copy
import re
import unicodedata
from datetime import UTC, datetime
from difflib import SequenceMatcher
from functools import lru_cache

from docx import Document
from docx.text.paragraph import Paragraph
from lxml import etree

from . import _xml
from ._xml import q
from .render import COLOR_CLASSES, _rgb, _style_color

AUTHOR = "ICF Studio"
# runs carrying these are kept as-is when a paragraph is rewritten (comment anchors, notes, fields)
_SPECIAL = {
    "commentReference", "footnoteReference", "endnoteReference", "drawing", "pict", "object",
    "fldChar", "instrText", "sym",
}  # fmt: skip
_TEXT = {"t", "tab", "br", "cr", "noBreakHyphen"}
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def norm(s: str) -> str:
    return " ".join(unicodedata.normalize("NFKC", s or "").split())


def _in_del(r) -> bool:
    return any(True for _ in r.iterancestors(q("del")))


def _style_name(p: Paragraph) -> str:
    try:
        return p.style.name if p.style is not None else ""
    except Exception:
        return ""


def list_paragraphs(path: str) -> list[dict]:
    """Non-empty paragraphs: {idx, text, section, heading}."""
    doc = Document(path)
    out: list[dict] = []
    section = ""
    for i, p_el in enumerate(doc.element.body.iter(q("p"))):
        text = _xml.para_text(p_el).strip()
        if not text:
            continue
        name = _style_name(Paragraph(p_el, doc._body)).lower()
        heading = name.startswith(("heading", "titre", "kop"))
        if heading:
            section = text[:120]
        out.append({"idx": i, "text": text, "section": section, "heading": heading})
    return out


# -- tracked-change rewrite ------------------------------------------------------------------


def _tokens(s: str) -> list[str]:
    return re.findall(r"\s+|[^\s]+", s)


def _segments(old: str, new: str) -> list[tuple[str, str]]:
    """Word-level ops [(eq|del|ins, text)] turning old into new."""
    a, b = _tokens(old), _tokens(new)
    sm = SequenceMatcher(None, a, b, autojunk=False)
    segs: list[tuple[str, str]] = []

    def add(op: str, text: str):
        if not text:
            return
        if segs and segs[-1][0] == op:
            segs[-1] = (op, segs[-1][1] + text)
        else:
            segs.append((op, text))

    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            add("eq", "".join(a[i1:i2]))
        else:
            add("del", "".join(a[i1:i2]))
            add("ins", "".join(b[j1:j2]))
    return segs


def _run(rpr, text: str, deleted: bool = False):
    r = etree.Element(q("r"))
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    tag = "delText" if deleted else "t"
    for k, part in enumerate(re.split(r"(\t|\n)", text)):
        if part == "\t":
            etree.SubElement(r, q("tab"))
        elif part == "\n":
            etree.SubElement(r, q("br"))
        elif part:
            t = etree.SubElement(r, q(tag))
            t.text = part
            t.set(_XML_SPACE, "preserve")
        del k
    return r


def _is_text_run(r) -> bool:
    kids = {etree.QName(c).localname for c in r}
    return bool(kids & _TEXT) and not (kids & _SPECIAL)


def _rewrite(p_el, old: str, new: str, author: str, date: str, ids) -> None:
    runs = [r for r in p_el.iter(q("r")) if not _in_del(r) and _is_text_run(r)]
    # formatting of every visible character, so unchanged/deleted text keeps its bold/colour
    char_rpr: list = []
    for r in runs:
        char_rpr.extend([r.find(q("rPr"))] * len(_xml.run_text(r)))
    old = "".join(_xml.run_text(r) for r in runs)
    default = next((x for x in char_rpr if x is not None), None)
    # anchor = the direct child of p that holds the first text run (keeps comment ranges around it)
    anchor = None
    if runs:
        anchor = runs[0]
        while anchor.getparent() is not p_el:
            anchor = anchor.getparent()

    def chunks(start: int, text: str):
        """Split text (old[start:...]) where the source run formatting changes."""
        i = 0
        while i < len(text):
            rpr = char_rpr[start + i] if start + i < len(char_rpr) else default
            j = i + 1
            while j < len(text) and start + j < len(char_rpr) and char_rpr[start + j] is rpr:
                j += 1
            yield rpr, text[i:j]
            i = j

    def tracked(kind: str):
        wrap = etree.Element(q(kind))
        wrap.set(q("id"), str(next(ids)))
        wrap.set(q("author"), author)
        wrap.set(q("date"), date)
        return wrap

    new_nodes = []
    pos_old = 0
    for op, text in _segments(old, new):
        if op == "ins":
            prev = char_rpr[pos_old - 1] if 0 < pos_old <= len(char_rpr) else default
            wrap = tracked("ins")
            wrap.append(_run(prev if prev is not None else default, text))
            new_nodes.append(wrap)
            continue
        if op == "eq":
            new_nodes.extend(_run(rpr, t) for rpr, t in chunks(pos_old, text))
        else:
            wrap = tracked("del")
            for rpr, t in chunks(pos_old, text):
                wrap.append(_run(rpr, t, deleted=True))
            new_nodes.append(wrap)
        pos_old += len(text)
    if anchor is not None:
        pos = list(p_el).index(anchor)
    else:
        ppr = p_el.find(q("pPr"))
        pos = 1 if ppr is not None else 0
    for n, node in enumerate(new_nodes):
        p_el.insert(pos + n, node)
    for r in runs:
        r.getparent().remove(r)


def apply_edits(src: str, out: str, edits: list[dict], author: str = AUTHOR) -> dict:
    """edits: [{idx, old, new}]. Writes tracked changes; returns {applied, skipped:[{idx,why}]}."""
    doc = Document(src)
    paras = list(doc.element.body.iter(q("p")))
    used = [int(x) for x in doc.element.xpath("//@w:id") if str(x).lstrip("-").isdigit()]
    counter = iter(range(max(used, default=0) + 1, 10**9))
    date = datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")
    applied, skipped = 0, []
    for e in edits:
        idx = int(e["idx"])
        if not 0 <= idx < len(paras):
            skipped.append({"idx": idx, "why": "paragraph not found"})
            continue
        p_el = paras[idx]
        current = _xml.para_text(p_el)
        if e.get("old") is not None and norm(current) != norm(e["old"]):
            skipped.append({"idx": idx, "why": "paragraph changed since the proposal"})
            continue
        if norm(current) == norm(e["new"]):
            continue
        _rewrite(p_el, current, e["new"], author, date, counter)
        applied += 1
    doc.save(out)
    return {"applied": applied, "skipped": skipped}


# -- mandatory-text guard --------------------------------------------------------------------


@lru_cache(maxsize=8)
def template_mandatory(path: str) -> tuple[str, ...]:
    """Paragraphs of a national template that are mostly purple (= mandatory wording)."""
    doc = Document(path)
    out: list[str] = []
    for p_el in doc.element.body.iter(q("p")):
        p = Paragraph(p_el, doc._body)
        try:
            pcolor = _style_color(p.style) if p.style is not None else None
        except Exception:
            pcolor = None
        purple = total = 0
        for r in p.runs:
            n = len(r.text.strip())
            if not n:
                continue
            try:
                color = _rgb(r.font) or _style_color(r.style) or pcolor
            except Exception:
                color = pcolor
            total += n
            if COLOR_CLASSES.get(color or "") == "purple":
                purple += n
        text = norm(p.text)
        if total and purple / total >= 0.6 and len(text) >= 40:
            out.append(text)
    return tuple(out)


def mandatory_match(text: str, mandatory: tuple[str, ...]) -> dict | None:
    """Best template match for a paragraph if it looks like mandatory wording (ratio >= 0.75)."""
    t = norm(text)
    if len(t) < 40:
        return None
    best, best_s = 0.0, None
    for m in mandatory:
        if abs(len(m) - len(t)) > max(len(m), len(t)) * 0.6:
            continue
        sm = SequenceMatcher(None, t, m, autojunk=False)
        if sm.real_quick_ratio() < best or sm.quick_ratio() < max(best, 0.75):
            continue
        r = sm.ratio()
        if r > best:
            best, best_s = r, m
    if best_s is None or best < 0.75:
        return None
    return {"score": round(best, 2), "template_text": best_s[:400]}
