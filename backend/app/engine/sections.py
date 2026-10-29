# ruff: noqa: E501
"""Section-level ICF drafting: a strong model edits whole template sections in context.

The national template is mostly black "proposal" text with blue placeholders and red guidance that
tells the author WHAT TO WRITE. Filling blue placeholders one by one (generate.py, legacy path)
loses that: the substantive trial text lives in the red guidance. Here each template chunk
(heading to heading) is shown to the model with its colour tags next to the full sponsor documents,
and the model returns paragraph edits: replace / delete / add paragraphs. Edits are applied to the
docx in place (formatting of the paragraph is kept); anything the model leaves alone falls through to
the legacy handling (red removed, blue flagged as 'Missing').
"""

from __future__ import annotations

import copy
import json
import logging
import os
import re
import time
from collections.abc import Callable
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait

from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.shared import RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from . import ai
from ._xml import q
from .render import COLOR_CLASSES, _style_color, is_toc_paragraph

log = logging.getLogger("icf.sections")

SPONSOR_CHARS = int(os.environ.get("ICF_SPONSOR_CHARS", "450000"))
BUDGET_S = int(os.environ.get("ICF_SECTION_BUDGET_S", "330"))
REVIEW_BUDGET_S = int(os.environ.get("ICF_REVIEW_BUDGET_S", "200"))
MAX_UNITS_PER_CHUNK = 45
LANG_NAMES = {"en": "English", "nl": "Dutch", "fr": "French"}
AUDIENCE = {
    "adult": "adult participants",
    "parent": "parents / legal guardians (addressed as 'you' about their child)",
    "adolescent": "adolescents aged 12-17 (clear, respectful plain language)",
    "child": "children aged 6-11 (very simple, friendly language)",
}
_TAG = {"P": "purple", "R": "red", "B": "blue", "G": "missing", "k": ""}

SYSTEM = """You are an expert regulatory medical writer who prepares Belgian informed consent forms (ICFs) for clinical trials inside the Belgian national (CT College) model ICF template. You get the sponsor's documents and ONE chunk of the template. You return the edits that turn that chunk into the final, trial-specific ICF text.

TEMPLATE COLOUR CODE (shown as tags in the chunk):
- <purple>text</purple> = MANDATORY text. Keep it word for word. Change only the minimum number of words if a sentence literally cannot be true for this trial (e.g. it speaks of an "IMP"/"study drug" while the trial has none). Never shorten, re-phrase or drop it otherwise.
- untagged text = proposal of the CT College. Keep it where it applies. Adapt it to this trial: replace the generic "IMP", "study drug", "disease/condition", "[number]" etc. by the real treatment, condition and values from the sponsor documents. Remove sentences/paragraphs that do not apply (e.g. guidance says "[if applicable]" / "[Choose]" / "[or]" and the trial lacks that element, or the paragraph is about something the trial does not have). When the template offers alternatives, keep only the one that matches the trial.
- <blue>text</blue> = placeholder to be replaced by trial-specific content from the sponsor documents (names, numbers, durations, lists, ...).
- <red>text</red> = instruction to the author (guidance). NEVER copy guidance into the output. FOLLOW it: it tells you what has to be written at that spot. Write that content from the sponsor documents (replace the guidance paragraph with the written text; use add_after for further paragraphs). If the instruction does not apply to this trial, delete the paragraph. Instructions about formatting (colours, boxes, printing) are simply deleted.
- <b>text</b> inside a tag = bold in the template.

CONTENT RULES
1. Use ONLY facts from the SPONSOR DOCUMENTS, the TRIAL BRIEF and the RESEARCHER NOTES. Never invent or guess facts, numbers, doses, dates, names, drugs, addresses. Copy numbers, doses, durations and names exactly. Everything you write must be in the TARGET LANGUAGE even if the documents are in another language (translate the facts).
2. If the RESEARCHER NOTES contain a "KNOWN FACTS" block, it is authoritative: use those values verbatim for the matching placeholders (site, contacts, insurer, ...). Otherwise site-specific values (hospital/site name and address, phone numbers, e-mail addresses of the site or investigator, investigator names, insurer name / policy number / address, CRO, payment vendors) must NEVER be invented. Use them only if literally present in the sponsor documents. Otherwise do not touch that paragraph (do not list it): the placeholder stays and is flagged automatically for the site.
3. When a trial-specific placeholder cannot be filled because the documents lack the information, also leave that paragraph unlisted (it is flagged automatically). Use {{MISSING: short description of the information needed}} only inside text you write when a passage required by the guidance cannot be written (e.g. the dose schedule) and no information exists. Do not use it for things the documents do contain: read the documents carefully, the information is usually there (possibly in another language, other wording or another document).
4. Blue text that is already a concrete value valid in every Belgian ICF (e.g. a fixed e-mail address or website of an authority) -> replace the paragraph with the same value in normal text.
5. Plain participant language, addressed to 'you': short sentences (< 15 words), short paragraphs, no jargon, explain medical terms. Use the fixed terms of the template: "trial", "trial staff", "investigator", "treating physicians" - but if the sponsor calls the research a (sub-)study keep the heading/terms consistent with the title of the trial.
6. Describe the real trial. If the trial has no investigational medicinal product (e.g. it compares radiation doses, a procedure or a strategy), re-write the IMP-oriented template sentences for the real intervention and delete those that cannot apply (extension trial with IMP, free access to IMP, study-drug side effect lists, compensation by vendor, extra samples, ...) - but keep every mandatory (purple) rule that still applies.
7. Text in [square brackets] (any colour), e.g. "[or if not possible: ...]", "[Choose]", "[if applicable]", "[number]", is an option or instruction: decide it, fill it from the documents or remove it. Never leave template brackets or "..." / "…" fillers in the final text of a paragraph you write.
8. Do not summarise away trial-specific detail that the guidance asks for: objectives, design (randomisation, probabilities, blinding), numbers of participants, schedule of visits/examinations, doses, duration, risks and side effects with frequencies, contraception/pregnancy rules and their duration, samples, benefits, alternatives, what happens after the trial. Write complete, correct and specific text.

OUTPUT: strict JSON {"edits": [ ... ]}. Only list paragraphs that change. Each edit:
  {"id": <paragraph id>, "op": "replace" | "delete" | "keep", "text": "<full new text of that paragraph>", "add_after": ["<new paragraph text>", ...]}
- "replace": "text" is the COMPLETE new paragraph (not just the changed part). Reproduce kept <purple> parts verbatim inside <purple>...</purple> tags; write bold as <b>...</b>; no other tags. Newlines are not allowed inside a paragraph.
- "delete": remove the paragraph (guidance, non-applicable text).
- "keep": leave the paragraph but insert "add_after" paragraphs after it.
- "add_after": optional new paragraphs inserted after this paragraph (they get the same paragraph format; for a bullet list put every bullet in its own string, without bullet characters).
Paragraphs you do not list stay unchanged (guidance-only paragraphs are removed automatically, unfilled blue placeholders are flagged automatically)."""


def _seg_runs(doc, p_el, p) -> list[Run]:
    from .generate import _runs

    return _runs(p_el, p)


def _run_color(run, para_color):
    from .generate import _run_color

    return _run_color(run, para_color)


def build_units(doc) -> list[dict]:
    """Every non-empty paragraph (tables included) with colour/bold tagged segments."""
    units = []
    for p_el in list(doc.element.body.iter(q("p"))):
        if is_toc_paragraph(p_el):
            continue
        p = Paragraph(p_el, doc._body)
        style = p.style.name if p.style is not None else ""
        pc = _style_color(p.style) if p.style is not None else None
        segs: list[list] = []  # [cls, bold, text]
        for r in _seg_runs(doc, p_el, p):
            t = r.text
            if not t:
                continue
            cls = COLOR_CLASSES.get(_run_color(r, pc) or "", "")
            cls = {"purple": "P", "red": "R", "blue": "B"}.get(cls, "k")
            if r.font.highlight_color == WD_COLOR_INDEX.YELLOW:
                cls = "G"  # gap already flagged for a human
            bold = bool(r.bold)
            if segs and segs[-1][0] == cls and segs[-1][1] == bold:
                segs[-1][2] += t
            else:
                segs.append([cls, bold, t])
        if not any(s[2].strip() for s in segs):
            continue
        low = style.lower()
        level = (
            1 if low.startswith(("heading 1", "title")) else 2 if low.startswith("heading 2") else 0
        )
        units.append(
            {
                "id": len(units),
                "el": p_el,
                "para": p,
                "style": style,
                "segs": segs,
                "level": level,
                "sub": low.startswith("subtitle"),
                "cell": p_el.getparent().tag == q("tc"),
                "listed": p_el.find(f"{q('pPr')}/{q('numPr')}") is not None,
                "protected": any(sg[0] == "G" for sg in segs),
            }
        )
    return units


def _fmt(u: dict) -> str:
    out = []
    for cls, bold, t in u["segs"]:
        if not t.strip():
            out.append(t)
            continue
        tag = _TAG[cls]
        s = f"<b>{t}</b>" if bold else t
        out.append(f"<{tag}>{s}</{tag}>" if tag else s)
    flags = []
    if u["level"]:
        flags.append(f"H{u['level']}")
    elif u["sub"]:
        flags.append("SUBHEADING")
    if u["cell"]:
        flags.append("table cell")
    if u["listed"]:
        flags.append("list item")
    head = f"[{u['id']}" + (f" | {', '.join(flags)}" if flags else "") + "]"
    return f"{head} {''.join(out).strip()}"


def make_chunks(units: list[dict]) -> list[list[dict]]:
    chunks: list[list[dict]] = []
    cur: list[dict] = []
    for u in units:
        if (u["level"] in (1, 2) or (u["sub"] and len(cur) >= MAX_UNITS_PER_CHUNK)) and cur:
            chunks.append(cur)
            cur = []
        elif len(cur) >= MAX_UNITS_PER_CHUNK * 2:
            chunks.append(cur)
            cur = []
        cur.append(u)
    if cur:
        chunks.append(cur)
    merged: list[list[dict]] = []
    for c in chunks:  # merge tiny chunks (a heading directly followed by a sub-heading) forward
        if merged and len(merged[-1]) < 4:
            merged[-1] = merged[-1] + c
        else:
            merged.append(c)
    return merged


def _sponsor_block(sponsor: str) -> str:
    sponsor = sponsor or ""
    if len(sponsor) <= SPONSOR_CHARS:
        return sponsor
    # keep the head of each document in equal measure when over budget
    docs = re.split(r"(?m)^(?==== )", sponsor)
    docs = [d for d in docs if d.strip()]
    per = SPONSOR_CHARS // max(len(docs), 1)
    return "\n".join(d[:per] for d in docs)


BRIEF_SYS = "You summarise clinical-trial documents precisely. Answer in strict JSON."
BRIEF_PROMPT = """From the sponsor documents below, write a TRIAL BRIEF as JSON with these keys (strings; use "unknown" when the documents do not say; write in {lang}):
trial_title (official title), lay_title, short_name, eu_ct_number, sponsor (name + address if present), disease_condition, population (age range, who is included), intervention (what participants receive/undergo in THIS trial incl. exact doses/schedules), has_imp (true/false: does the trial test an investigational medicinal product), comparator_or_arms (arms, randomisation ratio), n_participants (worldwide and Belgium), duration (per participant and total), procedures_and_visits (trial-specific examinations/visits), main_risks, contraception_pregnancy (rules and durations), samples (biological samples collected), costs_compensation (who pays, compensation), insurance, data_protection_specifics, terminology (how the documents call the research: study / trial / sub-study, and names of the treatments), notes_for_writer (anything unusual: e.g. sub-study of a platform trial refers to a main ICF).

SPONSOR DOCUMENTS:
{sponsor}
"""


def make_brief(sponsor: str, language: str) -> str:
    data, _ = ai.generate_json_models(
        BRIEF_PROMPT.format(
            lang=LANG_NAMES.get(language, "English"), sponsor=_sponsor_block(sponsor)
        ),
        BRIEF_SYS,
        [ai.fast_model()],
        thinking=0,
        timeout_ms=120_000,
        temperature=0,
    )
    if not isinstance(data, dict):
        return ""
    return json.dumps(data, ensure_ascii=False, indent=1)


def _outline(units: list[dict]) -> str:
    return "\n".join(
        ("  " if u["level"] == 2 else "") + "".join(s[2] for s in u["segs"]).strip()[:90]
        for u in units
        if u["level"] and not all(s[0] == "R" for s in u["segs"])
    )


def _chunk_prompt(chunk, *, sponsor, notes, brief, outline, language, icf_type, idx, total) -> str:
    return (
        f"TARGET LANGUAGE: {LANG_NAMES.get(language, 'English')} (write every word you add in this "
        f"language). AUDIENCE: {AUDIENCE.get(icf_type, 'adult participants')}.\n\n"
        f"=== SPONSOR DOCUMENTS ===\n{sponsor}\n\n"
        f"=== RESEARCHER NOTES ===\n{notes or '(none)'}\n\n"
        f"=== TRIAL BRIEF (machine summary of the documents; documents prevail) ===\n"
        f"{brief or '(not available)'}\n\n"
        f"=== OUTLINE OF THE WHOLE ICF (for cross-references and context) ===\n{outline}\n\n"
        f"=== TEMPLATE CHUNK {idx + 1}/{total} TO COMPLETE (paragraph ids in brackets) ===\n"
        + "\n".join(_fmt(u) for u in chunk)
        + "\n\nReturn the JSON edits for this chunk now."
    )


REVIEW_SYSTEM = """You are the quality reviewer of a Belgian informed consent form (ICF) for a clinical trial that was just drafted from the sponsor's documents inside the national (CT College) template. You get the sponsor documents and ONE chunk of the near-final draft. Fix only REAL problems and return paragraph edits; leave everything else untouched.

Tags: <purple> = mandatory template text (keep verbatim unless it is inapplicable to this trial), <b> = bold, <blue> = unfilled placeholder (leave it alone), <red> = left-over author guidance (delete it), <missing> = gap already flagged for a human (leave that paragraph alone). Untagged = normal text.

Look for, per paragraph:
1. Unsupported content: trial-specific statements (numbers, doses, durations, procedures, drugs, risks, rules, names) that are NOT in the sponsor documents/brief and are not generic ICF boilerplate. Delete the sentence or rewrite it to what the documents support; if a required passage cannot be supported write {{MISSING: what is needed}}.
2. Wrong or misleading content for THIS trial: text that describes elements the trial does not have (investigational medicinal product, study drug, compensation, vendors, extra visits/samples/examinations, extension trial, autologous product...) or contradicts the documents or other sentences in the chunk (e.g. "there are no trial-specific procedures" next to "trial-specific procedures are paid by the sponsor"). Delete or rewrite so the chunk is coherent.
3. Left-over template instructions, [brackets], "...", empty list introductions ("the following treatments:" with nothing after it), references to tables/lists/sections that do not exist, "(Ref. )" style empty references.
4. Wrong language (everything must be in the TARGET LANGUAGE) and garbled sentences.
Rules: never invent facts; never add new sections; keep plain participant language; keep mandatory text. Prefer small edits. If the chunk is fine return {"edits": []}.

OUTPUT: strict JSON {"edits": [{"id": <paragraph id>, "op": "replace" | "delete", "text": "<complete new paragraph text, <purple>/<b> tags allowed>"}]}. Paragraphs you do not list stay unchanged."""


def _run_chunk(chunk, ctx, idx, total):
    t0 = time.time()
    prompt = _chunk_prompt(chunk, idx=idx, total=total, **ctx["prompt"])
    data, spec = ai.generate_json_models(
        prompt,
        ctx.get("system", SYSTEM),
        ctx["models"],
        thinking=ctx["thinking"],
        timeout_ms=ctx["timeout_ms"],
        temperature=ctx["temperature"],
    )
    log.info(
        "section chunk %d/%d (%d paras) via %s in %.0fs",
        idx + 1,
        total,
        len(chunk),
        spec,
        time.time() - t0,
    )
    edits = data.get("edits") if isinstance(data, dict) else None
    return edits if isinstance(edits, list) else None


# -- applying edits ------------------------------------------------------------------------------

_TOKEN = re.compile(r"(</?(?:purple|b)>|\{\{\s*MISSING:.*?\}\})", re.S | re.I)
_STRAY = re.compile(r"</?(?:blue|red|k|i|u)>", re.I)


def _parse_inline(text: str) -> list[tuple[str, str, bool, bool]]:
    """-> [(kind, text, bold, purple)], kind 'text' | 'gap'."""
    text = _STRAY.sub("", text.replace("\r", "").replace("\n", " "))
    out = []
    bold = purple = False
    for part in _TOKEN.split(text):
        if not part:
            continue
        low = part.lower()
        if low == "<b>":
            bold = True
        elif low == "</b>":
            bold = False
        elif low == "<purple>":
            purple = True
        elif low == "</purple>":
            purple = False
        elif low.startswith("{{"):
            desc = re.sub(r"^\{\{\s*missing:\s*|\}\}$", "", part, flags=re.I).strip()
            out.append(("gap", desc, bold, purple))
        else:
            out.append(("text", part, bold, purple))
    return out


def _base_rpr(p_el):
    """rPr of the paragraph's first normal (black/purple) text run, cleaned of colour/bold/style."""
    cand = None
    for r in p_el.iter(q("r")):
        if any(True for _ in r.iterancestors(q("del"))):
            continue
        if not "".join(t.text or "" for t in r.iter(q("t"))).strip():
            continue
        rpr = r.find(q("rPr"))
        if cand is None:
            cand = rpr
        col = (
            rpr.find(q("color")).get(q("val"))
            if rpr is not None and rpr.find(q("color")) is not None
            else None
        )
        if COLOR_CLASSES.get((col or "").upper()) not in ("red", "blue"):
            cand = rpr
            break
    if cand is None:
        return None
    rpr = copy.deepcopy(cand)
    for tag in ("b", "bCs", "color", "highlight", "rStyle", "i", "iCs", "u", "shd", "strike"):
        for el in rpr.findall(q(tag)):
            rpr.remove(el)
    return rpr


def _write_paragraph(doc, p_el, text: str, base_rpr, style_colored: bool, gaps: list) -> int:
    """Rebuild the runs of p_el from tagged text; returns the number of gap markers."""
    for child in list(p_el):
        if child.tag != q("pPr"):
            p_el.remove(child)
    para = Paragraph(p_el, doc._body)
    n_gap = 0
    for kind, t, bold, purple in _parse_inline(text):
        if kind == "gap":
            t = f"[{t}]"
        run = para.add_run(t)
        if base_rpr is not None:
            run._r.insert(0, copy.deepcopy(base_rpr))
        if bold:
            run.font.bold = True
        if purple:
            run.font.color.rgb = RGBColor(0x70, 0x30, 0xA0)
        elif style_colored:
            run.font.color.rgb = RGBColor(0, 0, 0)
        if kind == "gap":
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            gaps.append((run, t))
            n_gap += 1
    return n_gap


def _clone_empty(p_el):
    new = copy.deepcopy(p_el)
    for child in list(new):
        if child.tag != q("pPr"):
            new.remove(child)
    ppr = new.find(q("pPr"))
    if ppr is not None:
        for tag in ("sectPr", "pPrChange"):
            for el in ppr.findall(q(tag)):
                ppr.remove(el)
    return new


def _wrap_fill(p_el, n: int) -> None:
    runs = [r for r in p_el if r.tag == q("r")]
    if not runs:
        return
    start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    bid = str(900000 + n)
    start.set(q("id"), bid)
    start.set(q("name"), f"icf_fill_{n}")
    end.set(q("id"), bid)
    runs[0].addprevious(start)
    runs[-1].addnext(end)


def _text_of(el) -> str:
    return "".join(t.text or "" for t in el.iter(q("t"))).strip()


def apply_edits(
    doc, units_by_id: dict[int, dict], edits: list, stats: dict, gaps: list, protect: bool = False
) -> None:
    from .generate import _drop_paragraph

    touched_tables = set()
    for e in edits:
        if not isinstance(e, dict):
            continue
        try:
            u = units_by_id.get(int(e.get("id")))
        except (TypeError, ValueError):
            u = None
        if u is None or u.get("done") or (protect and u.get("protected")):
            continue
        op = str(e.get("op") or "replace").lower()
        p_el = u["el"]
        u["done"] = True
        add = [a for a in (e.get("add_after") or []) if isinstance(a, str) and a.strip()]
        tbl = next(p_el.iterancestors(q("tbl")), None)
        if op == "delete" and not add:
            if tbl is not None:
                touched_tables.add(tbl)
            _drop_paragraph(p_el)
            stats["deleted"] += 1
            continue
        style_colored = COLOR_CLASSES.get(_style_color(u["para"].style) or "") in ("red", "blue")
        base = _base_rpr(p_el)
        text = e.get("text")
        if op == "replace" and isinstance(text, str) and text.strip():
            stats["gaps"] += _write_paragraph(doc, p_el, text, base, style_colored, gaps)
            _wrap_fill(p_el, stats["filled"])
            stats["filled"] += 1
        elif op == "delete":  # delete + add_after: replace the paragraph by the added ones
            if tbl is not None:
                touched_tables.add(tbl)
            first, add = add[0], add[1:]
            stats["gaps"] += _write_paragraph(doc, p_el, first, base, style_colored, gaps)
            _wrap_fill(p_el, stats["filled"])
            stats["filled"] += 1
        last = p_el
        for a in add:
            new = _clone_empty(p_el)
            last.addnext(new)
            last = new
            stats["gaps"] += _write_paragraph(doc, new, a, base, style_colored, gaps)
            _wrap_fill(new, stats["filled"])
            stats["filled"] += 1
    for tbl in touched_tables:  # drop rows/tables emptied by deletions
        if tbl.getparent() is None:
            continue
        for tr in list(tbl.findall(q("tr"))):
            if not _text_of(tr):
                tbl.remove(tr)
        if not tbl.findall(q("tr")) or not _text_of(tbl):
            parent = tbl.getparent()
            parent.remove(tbl)
            if parent.tag == q("tc") and parent.find(q("p")) is None:
                parent.append(OxmlElement("w:p"))


# -- driver --------------------------------------------------------------------------------------


def _run_stage(doc, units, ctx, report, offset, total, stats, gaps, budget_s, protect) -> int:
    """Run all chunks of `units` in parallel, apply the edits; returns the number of chunks ok."""
    chunks = [c for c in make_chunks(units) if protect is False or len(c) >= 3]
    if not chunks:
        return 0
    pool = ThreadPoolExecutor(max_workers=len(chunks))
    futs = {pool.submit(_run_chunk, c, ctx, i, len(chunks)): c for i, c in enumerate(chunks)}
    pending, done = set(futs), set()
    start = time.time()
    while pending:
        left = budget_s - (time.time() - start)
        if left <= 0:
            break
        fin, pending = wait(pending, timeout=left, return_when=FIRST_COMPLETED)
        done |= fin
        report("filling", offset + len(done), total)
    pool.shutdown(wait=False, cancel_futures=True)
    by_id = {u["id"]: u for u in units}
    ok = 0
    for f in sorted(done, key=lambda f: futs[f][0]["id"]):
        try:
            edits = f.result()
        except Exception as e:  # noqa: BLE001
            log.error("section chunk failed: %s", e)
            continue
        if edits is None:
            continue
        ok += 1
        apply_edits(doc, by_id, edits, stats, gaps, protect=protect)
    return ok


def fill_sections(
    doc,
    units: list[dict],
    *,
    sponsor_text: str,
    notes: str,
    language: str,
    icf_type: str,
    report: Callable[[str, int, int], None],
) -> dict:
    """Run the section pipeline on `doc` in place. Returns stats incl. gap runs for commenting.

    Stage 1 fills/edits every template chunk (strong model, full sponsor documents + trial brief).
    Stage 2 (ICF_REVIEW, default on) re-reads the resulting draft chunk by chunk and fixes
    unsupported or contradictory text and left-over template instructions.
    """
    stats = {"filled": 0, "deleted": 0, "gaps": 0, "chunks": 0, "chunks_ok": 0, "ai_used": False}
    gaps: list = []
    stats["gap_runs"] = gaps
    sponsor = _sponsor_block(sponsor_text)
    n_chunks = len(make_chunks(units))
    stats["chunks"] = n_chunks
    t0 = time.time()
    brief = make_brief(sponsor, language) if os.environ.get("ICF_BRIEF", "1") != "0" else ""
    log.info(
        "section mode: %d units, %d chunks, brief %d chars (%.0fs)",
        len(units),
        n_chunks,
        len(brief),
        time.time() - t0,
    )
    review = os.environ.get("ICF_REVIEW", "1") != "0"
    total = n_chunks * (2 if review else 1)
    base = {
        "sponsor": sponsor,
        "notes": notes,
        "brief": brief,
        "outline": _outline(units),
        "language": language,
        "icf_type": icf_type,
    }
    ctx = {
        "prompt": base,
        "models": ai.fill_models(),
        "thinking": os.environ.get("ICF_FILL_THINKING") or None,
        "timeout_ms": int(os.environ.get("ICF_FILL_TIMEOUT_S", "240")) * 1000,
        "temperature": float(os.environ.get("ICF_FILL_TEMPERATURE", "0.2")),
    }
    report("filling", 0, total)
    stats["chunks_ok"] = _run_stage(doc, units, ctx, report, 0, total, stats, gaps, BUDGET_S, False)
    stats["ai_used"] = stats["chunks_ok"] > 0
    log.info(
        "section stage 1: %d/%d chunks ok, %d paragraphs written, %d deleted, %d gaps (%.0fs)",
        stats["chunks_ok"],
        n_chunks,
        stats["filled"],
        stats["deleted"],
        stats["gaps"],
        time.time() - t0,
    )
    if review and stats["chunks_ok"] == n_chunks:
        units2 = build_units(doc)
        rctx = dict(
            ctx,
            system=REVIEW_SYSTEM,
            prompt=dict(base, outline=_outline(units2)),
            thinking=os.environ.get("ICF_REVIEW_THINKING") or ctx["thinking"],
        )
        ok = _run_stage(
            doc, units2, rctx, report, n_chunks, total, stats, gaps, REVIEW_BUDGET_S, True
        )
        log.info("section stage 2 (review): %d chunks ok (%.0fs)", ok, time.time() - t0)
    report("filling", total, total)
    return stats
