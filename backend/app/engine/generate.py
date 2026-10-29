"""Generate an ICF draft from a Belgian template: strip guidance, fill/flag placeholders."""

import json
import logging
import os
import time
from collections.abc import Callable
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from . import ai
from ._xml import q
from .render import COLOR_CLASSES, _rgb, _style_color, is_toc_paragraph

log = logging.getLogger("icf.generate")

AI_BUDGET_S = 150
BATCH_SIZE = 6
SPONSOR_CHARS = 40_000
AUTHOR = "ICF Studio"
LANG_NAMES = {"en": "English", "nl": "Dutch", "fr": "French"}
ICF_AUDIENCE = {
    "adult": "adult participants",
    "parent": "parents / legal guardians (addressed as 'you' about their child)",
    "adolescent": "adolescents aged 12-17 (clear, respectful plain language)",
    "child": "children aged 6-11 (very simple, friendly language)",
}

SYSTEM = (
    "You complete placeholders of an informed consent form (ICF) for a clinical trial in children "
    "with cancer. You may ONLY use facts that are explicitly present in the SPONSOR DOCUMENTS or "
    "RESEARCHER NOTES. Never invent or guess facts, numbers, names, drugs, addresses or dates. "
    "If the provided text does not contain the information for a placeholder, answer exactly "
    "MISSING for it. Otherwise write only the replacement text (no brackets, no explanations), "
    "ALWAYS written in the requested target language, even when the sponsor documents or notes are "
    "in another language (translate the facts), at a level the participant understands."
)


def _run_color(run: Run, para_color: str | None) -> str | None:
    try:
        return _rgb(run.font) or _style_color(run.style) or para_color
    except Exception:
        return para_color


def _runs(p_el, p: Paragraph) -> list[Run]:
    return [
        Run(r, p) for r in p_el.iter(q("r")) if not any(True for _ in r.iterancestors(q("del")))
    ]


def _has_sect(p_el) -> bool:
    return p_el.find(f"{q('pPr')}/{q('sectPr')}") is not None


def _drop_paragraph(p_el):
    parent = p_el.getparent()
    if _has_sect(p_el) or (parent.tag == q("tc") and len(parent.findall(q("p"))) <= 1):
        for r in list(p_el.iter(q("r"))):
            r.getparent().remove(r)
        return
    parent.remove(p_el)


def _strip_preamble(doc) -> int:
    """Remove the template's own instruction pages before the first red heading."""
    body = doc.element.body
    children = list(body.iterchildren())
    seen_guidance = False
    cut = None
    for i, el in enumerate(children):
        if el.tag != q("p"):
            continue
        p = Paragraph(el, doc._body)
        name = p.style.name if p.style is not None else ""
        text = p.text.strip()
        if name.lower().startswith("heading") and text.lower() == "guidance":
            seen_guidance = True
        elif seen_guidance and name.lower().startswith("heading") and text:
            runs = [r for r in _runs(el, p) if r.text.strip()]
            if runs and all(_run_color(r, _style_color(p.style)) == "FF0000" for r in runs):
                cut = i
                break
    if cut is None:
        return 0
    removed = 0
    for el in children[:cut]:
        if el.tag == q("p") and not _has_sect(el):
            if Paragraph(el, doc._body).text.strip():
                removed += 1
            body.remove(el)
        elif el.tag == q("tbl"):
            body.remove(el)
    return removed


def _collect_placeholders(doc) -> tuple[list[dict], int]:
    """Remove red guidance; return blue placeholder groups and removed paragraph count."""
    removed = 0
    placeholders: list[dict] = []
    section = ""
    for p_el in list(doc.element.body.iter(q("p"))):
        if is_toc_paragraph(p_el):  # never fill, strip or flag the table of contents
            continue
        p = Paragraph(p_el, doc._body)
        style_name = p.style.name if p.style is not None else ""
        para_color = _style_color(p.style) if p.style is not None else None
        runs = _runs(p_el, p)
        text_runs = [r for r in runs if r.text]
        reds = [r for r in text_runs if _run_color(r, para_color) == "FF0000"]
        if reds:
            keep = [r for r in text_runs if r not in reds]
            for r in reds:
                r._r.getparent().remove(r._r)
            if not any(r.text.strip() for r in keep):
                removed += 1
                _drop_paragraph(p_el)
                continue
            runs = [r for r in runs if r not in reds]
        if style_name.lower().startswith(("heading", "titre")) and p.text.strip():
            section = p.text.strip()
        group: list[Run] = []
        pending: list[Run] = []  # whitespace-only runs after the current group

        def flush(group=group, p=p, section=section):
            if group:
                text = "".join(r.text for r in group).strip()
                if text:
                    placeholders.append(
                        {"runs": list(group), "text": text, "section": section, "para": p}
                    )
                group.clear()

        for r in runs:
            if not r.text:
                continue
            if COLOR_CLASSES.get(_run_color(r, para_color) or "") == "blue":
                group.extend(pending)
                pending.clear()
                group.append(r)
            elif group and not r.text.strip():
                pending.append(r)
            else:
                pending.clear()
                flush()
        pending.clear()
        flush()
    for i, ph in enumerate(placeholders):
        ph["id"] = f"p{i}"
    return placeholders, removed


def _context(ph: dict) -> str:
    """Paragraph text with this placeholder marked."""
    first, last = ph["runs"][0]._r, ph["runs"][-1]._r
    parts, inside = [], False
    for r in ph["para"]._p.iter(q("r")):
        t = Run(r, ph["para"]).text
        if r is first:
            parts.append("<<")
            inside = True
        parts.append(t)
        if r is last and inside:
            parts.append(">>")
            inside = False
    return "".join(parts)[:600]


def _fill_batch(
    batch: list[dict], sponsor: str, notes: str, language: str, icf_type: str, thinking: int
):
    items = [
        {
            "id": ph["id"],
            "section": ph["section"],
            "context": _context(ph),
            "placeholder": ph["text"],
        }
        for ph in batch
    ]
    prompt = (
        f"Target language: {LANG_NAMES.get(language, 'English')} - write every replacement in this "
        "language even if the sponsor documents are in another language. Audience: "
        f"{ICF_AUDIENCE.get(icf_type, 'adult participants')}.\n\n"
        f"SPONSOR DOCUMENTS:\n{sponsor}\n\nRESEARCHER NOTES:\n{notes or '(none)'}\n\n"
        "PLACEHOLDERS (JSON). In each 'context' the placeholder to replace is wrapped in << >>. "
        "Your answer replaces exactly the wrapped text.\n"
        f"{json.dumps(items, ensure_ascii=False)}\n\n"
        'Return a JSON object mapping each placeholder id to its replacement text, or to "MISSING" '
        "when the sponsor documents and notes do not contain that information."
    )
    data = ai.generate_json(prompt, SYSTEM, thinking_budget=thinking)
    return data if isinstance(data, dict) else None


def _valid(ans) -> str | None:
    if not isinstance(ans, str):
        return None
    ans = ans.strip()
    if not ans or ans.upper().startswith("MISSING"):
        return None
    return ans


def _mark_fill(r_el, n: int) -> None:
    """Wrap a filled run in a bookmark `icf_fill_N` so the viewer can highlight it (no comment)."""
    from docx.oxml import OxmlElement

    start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    bid = str(900000 + n)
    start.set(q("id"), bid)
    start.set(q("name"), f"icf_fill_{n}")
    end.set(q("id"), bid)
    r_el.addprevious(start)
    r_el.addnext(end)


def _refresh_toc_on_open(doc) -> None:
    """Ask Word to refresh fields (TOC page numbers / removed guidance entries) when opening."""
    from docx.oxml import OxmlElement

    settings = doc.settings.element
    if settings.find(q("updateFields")) is not None:
        return
    el = OxmlElement("w:updateFields")
    el.set(q("val"), "true")
    later = {  # schema order: updateFields precedes these
        "hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr",
        "attachedSchema", "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats",
        "doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown",
        "smartTagType", "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol",
        "listSeparator",
    }  # fmt: skip
    nxt = next((c for c in settings if c.tag.rsplit("}", 1)[-1] in later), None)
    if nxt is not None:
        nxt.addprevious(el)
    else:
        settings.append(el)


def generate_icf(
    template_path: str,
    out_path: str,
    *,
    sponsor_text: str,
    notes: str,
    language: str,
    icf_type: str,
    progress: Callable[[str, int, int], None] | None = None,
    thinking_budget: int = 1024,
) -> dict:
    def report(step: str, done: int, total: int) -> None:
        if progress:
            try:
                progress(step, done, total)
            except Exception:  # noqa: BLE001
                log.exception("progress callback failed")

    report("reading", 0, 1)
    doc = Document(template_path)
    removed = _strip_preamble(doc)
    ai_on = bool(ai.ai_status()["enabled"])
    mode = os.environ.get("ICF_GEN_MODE", "sections")
    sec: dict | None = None
    if mode == "sections" and ai_on:
        from . import sections

        units = sections.build_units(doc)
        report("reading", 1, 1)
        sec = sections.fill_sections(
            doc,
            units,
            sponsor_text=sponsor_text or "",
            notes=notes,
            language=language,
            icf_type=icf_type,
            report=report,
        )
    placeholders, r2 = _collect_placeholders(doc)  # strips leftover guidance, collects blue left
    removed += r2
    if sec is None:
        report("reading", 1, 1)
    answers: dict[str, str] = {}
    ai_used = bool(sec and sec["ai_used"])
    # legacy placeholder-by-placeholder filling: whole pipeline in "placeholders" mode, and as the
    # fallback when the section pipeline produced nothing (e.g. only a weak model is reachable)
    if placeholders and ai_on and (sec is None or sec["chunks_ok"] < sec["chunks"]):
        sponsor = (sponsor_text or "")[:SPONSOR_CHARS]
        n = BATCH_SIZE
        batches = [placeholders[i : i + n] for i in range(0, len(placeholders), n)]
        start = time.time()
        pool = ThreadPoolExecutor(
            max_workers=max(1, len(batches))
        )  # ai._slots caps in-flight calls
        futs = {
            pool.submit(_fill_batch, b, sponsor, notes, language, icf_type, thinking_budget): b
            for b in batches
        }
        report("filling", 0, len(batches))
        done: set = set()
        pending = set(futs)
        while pending:
            left = AI_BUDGET_S - (time.time() - start)
            if left <= 0:
                break
            fin, pending = wait(pending, timeout=left, return_when=FIRST_COMPLETED)
            done |= fin
            report("filling", len(done), len(batches))
        pool.shutdown(wait=False, cancel_futures=True)
        log.info("generate: %d/%d batches in %.0fs", len(done), len(batches), time.time() - start)
        for f in done:
            try:
                data = f.result()
            except Exception as e:  # noqa: BLE001
                log.error("batch failed: %s", e)
                continue
            if data is None:
                continue
            ai_used = True
            for ph in futs[f]:
                ans = _valid(data.get(ph["id"]))
                if ans:
                    answers[ph["id"]] = ans

    filled = sec["filled"] if sec else 0
    flagged = 0
    total_steps = len(placeholders) + (len(sec["gap_runs"]) if sec else 0)
    report("comments", 0, total_steps)
    k = 0
    for run, desc in sec["gap_runs"] if sec else []:  # gaps written by the section pipeline
        shown = desc if len(desc) <= 200 else desc[:200] + "..."
        doc.add_comment(
            [run],
            text=f"Missing: {shown.strip('[]')} - please add this information",
            author=AUTHOR,
            initials="ICF",
        )
        flagged += 1
        k += 1
        report("comments", k, total_steps)
    for n, ph in enumerate(placeholders, k + 1):
        runs = ph["runs"]
        ans = answers.get(ph["id"])
        if ans:
            first = runs[0]
            first.text = ans
            for r in runs[1:]:
                r._r.getparent().remove(r._r)
            first.font.color.rgb = None
            if COLOR_CLASSES.get(_style_color(first.style) or "") == "blue":
                from docx.shared import RGBColor

                first.font.color.rgb = RGBColor(0, 0, 0)
            _mark_fill(first._r, filled)  # quiet marker, no Word comment
            filled += 1
        else:
            for r in runs:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            shown = ph["text"] if len(ph["text"]) <= 200 else ph["text"][:200] + "..."
            doc.add_comment(
                runs,
                text=f"Missing: {shown} - please add this information",
                author=AUTHOR,
                initials="ICF",
            )
            flagged += 1
        report("comments", n, total_steps)
    report("saving", 0, 1)
    _refresh_toc_on_open(doc)
    doc.save(out_path)
    return {
        "placeholders_total": len(placeholders),
        "filled": filled,
        "flagged": flagged,
        "removed_guidance_paragraphs": removed,
        "ai_used": ai_used,
    }
