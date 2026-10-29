"""Colour-aware DOCX -> HTML renderer with Word comment ranges and tracked changes."""

import re
from html import escape

from docx import Document as load_docx
from docx.text.run import Run

from ._xml import q

COLOR_CLASSES = {"7030A0": "purple", "FF0000": "red", "0000FF": "blue"}
_INLINE_WRAPPERS = {"hyperlink", "smartTag", "fldSimple", "sdt", "sdtContent", "customXml"}


def _rgb(font) -> str | None:
    try:
        rgb = font.color.rgb if font.color is not None else None
    except Exception:
        return None
    return str(rgb).upper() if rgb is not None else None


def _style_color(style) -> str | None:
    while style is not None:
        c = _rgb(style.font)
        if c:
            return c
        style = style.base_style
    return None


def _slug(text: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return s[:40] or "section"


def _heading_level(style_name: str) -> int | None:
    low = style_name.lower()
    if not low.startswith(("heading", "titre")):
        return None
    m = re.search(r"(\d+)\s*$", style_name)
    return min(max(int(m.group(1)), 1), 4) if m else 1


class _Renderer:
    def __init__(self, doc):
        self.doc = doc
        self.headings: list[dict] = []
        self.paragraphs = 0
        self.tables = 0
        self.colors = {"purple": 0, "red": 0, "blue": 0, "black": 0}
        self.used_anchors: set[str] = set()
        self.comment_ids: list[str] = []
        self.fills = 0
        self.fill_ids: set[str] = set()  # open icf_fill bookmarks
        self.active: list[str] = []  # comment ranges open at this point in the document
        self._num_cache: dict[tuple[str, str], str] = {}

    # -- inline -------------------------------------------------------------
    def _run_html(self, r_el, p, para_color) -> tuple[str, str]:
        """Return (html, plain visible text) for one w:r."""
        in_del = any(True for _ in r_el.iterancestors(q("del")))
        text = ""
        for c in r_el:
            tag = c.tag.rsplit("}", 1)[-1]
            if tag == "t" or (tag == "delText" and in_del):
                text += c.text or ""
            elif tag == "tab":
                text += "\t"
            elif tag in ("br", "cr"):
                text += "\n"
            elif tag == "noBreakHyphen":
                text += "-"
        if not text:
            return "", ""
        run = Run(r_el, p)
        try:
            color = _rgb(run.font) or _style_color(run.style) or para_color
        except Exception:
            color = para_color
        cls = COLOR_CLASSES.get(color or "")
        if not in_del:
            self.colors[cls or "black"] += len(text.split())
        h = escape(text).replace("\n", "<br>").replace("\t", "&emsp;")
        if run.bold:
            h = f"<strong>{h}</strong>"
        if run.italic:
            h = f"<em>{h}</em>"
        if run.underline:
            h = f"<u>{h}</u>"
        if cls:
            h = f'<span class="c-{cls}">{h}</span>'
        change = r_el.find(f"{q('rPr')}/{q('rPrChange')}")
        if change is not None:  # tracked formatting change (e.g. made bold)
            author = escape(change.get(q("author")) or "", quote=True)
            h = f'<span class="tc-fmt" data-author="{author}">{h}</span>'
        return h, ("" if in_del else text)

    def _walk_inline(self, el, p, para_color, pieces: list, plain: list):
        for child in el:
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "commentRangeStart":
                cid = child.get(q("id"))
                if cid is not None and cid not in self.active:
                    self.active.append(cid)
                    if cid not in self.comment_ids:
                        self.comment_ids.append(cid)
            elif tag == "commentRangeEnd":
                cid = child.get(q("id"))
                if cid in self.active:
                    self.active.remove(cid)
            elif tag == "bookmarkStart":
                if (child.get(q("name")) or "").startswith("icf_fill"):
                    self.fill_ids.add(child.get(q("id")))
                    self.fills += 1
            elif tag == "bookmarkEnd":
                self.fill_ids.discard(child.get(q("id")))
            elif tag == "r":
                h, t = self._run_html(child, p, para_color)
                if h and self.fill_ids:
                    h = f'<mark class="fill">{h}</mark>'
                if h:
                    pieces.append((tuple(self.active), h, self._tc(child)))
                    plain.append(t)
            elif tag in _INLINE_WRAPPERS:
                self._walk_inline(child, p, para_color, pieces, plain)
            elif tag in ("ins", "del"):
                self._walk_inline(child, p, para_color, pieces, plain)

    @staticmethod
    def _tc(r_el) -> tuple | None:
        """Innermost tracked-change wrapper of a run: (kind, author, date)."""
        for a in r_el.iterancestors():
            tag = a.tag.rsplit("}", 1)[-1]
            if tag in ("ins", "del"):
                return (tag, a.get(q("author")) or "", a.get(q("date")) or "")
            if tag == "p":
                break
        return None

    def runs_html(self, p_el, p, para_color) -> tuple[str, str]:
        pieces: list = []
        plain: list[str] = []
        self._walk_inline(p_el, p, para_color, pieces, plain)
        out: list[str] = []
        i = 0
        while i < len(pieces):
            tc = pieces[i][2]
            j = i
            seg: list[str] = []
            while j < len(pieces) and pieces[j][2] == tc:
                seg.append(pieces[j])
                j += 1
            inner = self._wrap_cmt(seg)
            if tc:
                kind, author, date = tc
                cls = "tc-ins" if kind == "ins" else "tc-del"
                inner = (
                    f'<{kind} class="{cls}" data-author="{escape(author, quote=True)}" '
                    f'data-date="{escape(date, quote=True)}">{inner}</{kind}>'
                )
            out.append(inner)
            i = j
        return "".join(out), "".join(plain)

    @staticmethod
    def _wrap_cmt(seg: list) -> str:
        out: list[str] = []
        i = 0
        while i < len(seg):
            ids = seg[i][0]
            j = i
            buf = []
            while j < len(seg) and seg[j][0] == ids:
                buf.append(seg[j][1])
                j += 1
            h = "".join(buf)
            if ids:
                h = f'<span class="cmt" data-cid="{escape(" ".join(ids), quote=True)}">{h}</span>'
            out.append(h)
            i = j
        return "".join(out)

    # -- lists --------------------------------------------------------------
    def _list_kind(self, p) -> str | None:
        name = p.style.name.lower() if p.style is not None else ""
        ppr = p._p.pPr
        num_pr = ppr.numPr if ppr is not None else None
        if num_pr is not None and num_pr.numId is not None and num_pr.numId.val != 0:
            ilvl = num_pr.ilvl.val if num_pr.ilvl is not None else 0
            return self._num_format(str(num_pr.numId.val), str(ilvl))
        if "list bullet" in name:
            return "ul"
        if "list number" in name:
            return "ol"
        return None

    def _num_format(self, num_id: str, ilvl: str) -> str:
        key = (num_id, ilvl)
        if key in self._num_cache:
            return self._num_cache[key]
        kind = "ul"
        try:
            root = self.doc.part.numbering_part.element
            abstract_id = None
            for num in root.findall(q("num")):
                if num.get(q("numId")) == num_id:
                    a = num.find(q("abstractNumId"))
                    abstract_id = a.get(q("val")) if a is not None else None
            for a in root.findall(q("abstractNum")):
                if a.get(q("abstractNumId")) != abstract_id:
                    continue
                for lvl in a.findall(q("lvl")):
                    if lvl.get(q("ilvl")) == ilvl:
                        fmt = lvl.find(q("numFmt"))
                        val = fmt.get(q("val")) if fmt is not None else "bullet"
                        kind = "ul" if val in ("bullet", "none") else "ol"
        except Exception:
            kind = "ul"
        self._num_cache[key] = kind
        return kind

    # -- blocks -------------------------------------------------------------
    def paragraph(self, p_el) -> tuple[str, str | None]:
        from docx.text.paragraph import Paragraph

        p = Paragraph(p_el, self.doc._body)
        style_name = p.style.name if p.style is not None else ""
        para_color = _style_color(p.style) if p.style is not None else None
        inner, plain = self.runs_html(p_el, p, para_color)
        if not inner:
            return "", None
        if plain.strip():
            self.paragraphs += 1
        level = _heading_level(style_name)
        if level is not None and plain.strip():
            base = _slug(plain)
            anchor, n = base, 2
            while anchor in self.used_anchors:
                anchor, n = f"{base}-{n}", n + 1
            self.used_anchors.add(anchor)
            self.headings.append({"level": level, "text": plain.strip(), "anchor": anchor})
            return f'<h{level} id="{anchor}">{inner}</h{level}>', None
        kind = self._list_kind(p)
        if kind:
            return f"<li>{inner}</li>", kind
        return f"<p>{inner}</p>", None

    def table(self, tbl) -> str:
        self.tables += 1
        rows = []
        for tr in tbl.findall(q("tr")):
            cells = []
            for tc in tr.findall(q("tc")):
                span = 1
                gs = tc.find(f"{q('tcPr')}/{q('gridSpan')}")
                if gs is not None:
                    try:
                        span = int(gs.get(q("val")))
                    except (TypeError, ValueError):
                        span = 1
                attr = f' colspan="{span}"' if span > 1 else ""
                cells.append(f"<td{attr}>{self.blocks(tc)}</td>")
            rows.append(f"<tr>{''.join(cells)}</tr>")
        return f'<table class="docx-table"><tbody>{"".join(rows)}</tbody></table>'

    def blocks(self, container) -> str:
        out: list[str] = []
        open_list: str | None = None
        for child in container.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            items: list[tuple[str, str | None]] = []
            if tag == "p":
                items.append(self.paragraph(child))
            elif tag == "tbl":
                items.append((self.table(child), None))
            elif tag == "sdt":
                content = child.find(q("sdtContent"))
                if content is not None:
                    items.append((self.blocks(content), None))
            for html, kind in items:
                if not html:
                    continue
                if kind != open_list:
                    if open_list:
                        out.append(f"</{open_list}>")
                    if kind:
                        out.append(f"<{kind}>")
                    open_list = kind
                out.append(html)
        if open_list:
            out.append(f"</{open_list}>")
        return "\n".join(out)


def render_docx(path: str) -> dict:
    doc = load_docx(path)
    r = _Renderer(doc)
    html = r.blocks(doc.element.body)
    return {
        "html": html,
        "headings": r.headings,
        "stats": {
            "paragraphs": r.paragraphs,
            "words": sum(r.colors.values()),
            "tables": r.tables,
            "colors": r.colors,
        },
        "comment_ids": r.comment_ids,
        "fill_count": r.fills,
    }
