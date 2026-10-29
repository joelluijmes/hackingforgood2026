"""Colour-aware DOCX -> HTML renderer. All text is escaped with html.escape."""

import re
from html import escape
from io import BytesIO

from docx import Document as load_docx
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

COLOR_CLASSES = {"7030A0": "purple", "FF0000": "red", "0000FF": "blue"}
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _rgb(font) -> str | None:
    try:
        rgb = font.color.rgb if font.color is not None else None
    except Exception:
        return None
    return str(rgb).upper() if rgb is not None else None


def _style_color(style) -> str | None:
    while style is not None:
        c = _rgb(style.font)
        if c:
            return c
        style = style.base_style
    return None


def _slug(text: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return s[:40] or "section"


def _heading_level(style_name: str) -> int | None:
    low = style_name.lower()
    if not low.startswith(("heading", "titre")):
        return None
    m = re.search(r"(\d+)\s*$", style_name)
    return min(max(int(m.group(1)), 1), 4) if m else 1


class _Renderer:
    def __init__(self, doc):
        self.doc = doc
        self.headings: list[dict] = []
        self.paragraphs = 0
        self.tables = 0
        self.colors = {"purple": 0, "red": 0, "blue": 0, "black": 0}
        self.used_anchors: set[str] = set()
        self._num_cache: dict[tuple[str, str], str] = {}

    # -- inline -------------------------------------------------------------
    def runs_html(self, p: Paragraph, para_color: str | None) -> tuple[str, int]:
        out: list[str] = []
        total = 0
        try:
            items = list(p.iter_inner_content())
        except AttributeError:
            items = list(p.runs)
        runs = []
        for it in items:
            runs.extend(it.runs if hasattr(it, "runs") else [it])
        for r in runs:
            text = r.text
            if not text:
                continue
            color = _rgb(r.font) or _style_color(r.style) or para_color
            cls = COLOR_CLASSES.get(color or "")
            words = len(text.split())
            total += words
            self.colors[cls or "black"] += words
            h = escape(text).replace("\n", "<br>").replace("\t", "&emsp;")
            if r.bold:
                h = f"<strong>{h}</strong>"
            if r.italic:
                h = f"<em>{h}</em>"
            if r.underline:
                h = f"<u>{h}</u>"
            out.append(f'<span class="c-{cls}">{h}</span>' if cls else h)
        return "".join(out), total

    # -- lists --------------------------------------------------------------
    def _list_kind(self, p: Paragraph) -> str | None:
        name = p.style.name.lower() if p.style is not None else ""
        ppr = p._p.pPr
        num_pr = ppr.numPr if ppr is not None else None
        if num_pr is not None and num_pr.numId is not None and num_pr.numId.val != 0:
            ilvl = num_pr.ilvl.val if num_pr.ilvl is not None else 0
            return self._num_format(str(num_pr.numId.val), str(ilvl))
        if "list bullet" in name:
            return "ul"
        if "list number" in name:
            return "ol"
        return None

    def _num_format(self, num_id: str, ilvl: str) -> str:
        key = (num_id, ilvl)
        if key in self._num_cache:
            return self._num_cache[key]
        kind = "ul"
        try:
            root = self.doc.part.numbering_part.element
            abstract_id = None
            for num in root.findall(f"{W_NS}num"):
                if num.get(f"{W_NS}numId") == num_id:
                    a = num.find(f"{W_NS}abstractNumId")
                    abstract_id = a.get(f"{W_NS}val") if a is not None else None
            for a in root.findall(f"{W_NS}abstractNum"):
                if a.get(f"{W_NS}abstractNumId") != abstract_id:
                    continue
                for lvl in a.findall(f"{W_NS}lvl"):
                    if lvl.get(f"{W_NS}ilvl") == ilvl:
                        fmt = lvl.find(f"{W_NS}numFmt")
                        val = fmt.get(f"{W_NS}val") if fmt is not None else "bullet"
                        kind = "ul" if val in ("bullet", "none") else "ol"
        except Exception:
            kind = "ul"
        self._num_cache[key] = kind
        return kind

    # -- blocks -------------------------------------------------------------
    def paragraph(self, p: Paragraph) -> tuple[str, str | None]:
        """Return (html, list_kind). List items come back as <li> for the caller to group."""
        style_name = p.style.name if p.style is not None else ""
        para_color = _style_color(p.style) if p.style is not None else None
        inner, words = self.runs_html(p, para_color)
        if not p.text.strip():
            return "", None
        self.paragraphs += 1
        level = _heading_level(style_name)
        if level is not None:
            base = _slug(p.text)
            anchor, n = base, 2
            while anchor in self.used_anchors:
                anchor, n = f"{base}-{n}", n + 1
            self.used_anchors.add(anchor)
            self.headings.append({"level": level, "text": p.text.strip(), "anchor": anchor})
            return f'<h{level} id="{anchor}">{inner}</h{level}>', None
        kind = self._list_kind(p)
        if kind:
            return f"<li>{inner}</li>", kind
        return f"<p>{inner}</p>", None

    def table(self, t: Table) -> str:
        self.tables += 1
        rows = []
        for tr in t._tbl.tr_lst:
            cells = []
            for tc in tr.tc_lst:
                cell = _Cell(tc, t)
                span = tc.grid_span
                parts = []
                for cp in cell.paragraphs:
                    inner, _ = self.runs_html(cp, _style_color(cp.style))
                    if cp.text.strip():
                        self.paragraphs += 1
                        parts.append(f"<p>{inner}</p>")
                attr = f' colspan="{span}"' if span > 1 else ""
                cells.append(f"<td{attr}>{''.join(parts)}</td>")
            rows.append(f"<tr>{''.join(cells)}</tr>")
        return f'<table class="docx-table"><tbody>{"".join(rows)}</tbody></table>'

    def render(self) -> str:
        out: list[str] = []
        open_list: str | None = None
        for child in self.doc.element.body.iterchildren():
            tag = child.tag.removeprefix(W_NS)
            if tag == "p":
                html, kind = self.paragraph(Paragraph(child, self.doc._body))
                if kind != open_list:
                    if open_list:
                        out.append(f"</{open_list}>")
                    if kind:
                        out.append(f"<{kind}>")
                    open_list = kind
                if html:
                    out.append(html)
            elif tag == "tbl":
                if open_list:
                    out.append(f"</{open_list}>")
                    open_list = None
                out.append(self.table(Table(child, self.doc._body)))
        if open_list:
            out.append(f"</{open_list}>")
        return "\n".join(out)


def render_docx(data: bytes) -> dict:
    doc = load_docx(BytesIO(data))
    r = _Renderer(doc)
    html = r.render()
    words = sum(r.colors.values())
    return {
        "html": html,
        "headings": r.headings,
        "stats": {
            "paragraphs": r.paragraphs,
            "words": words,
            "tables": r.tables,
            "colors": r.colors,
        },
    }
